"""Trace Excel precedents from extracted cells without pretending to recalculate Excel."""
from __future__ import annotations

import re
from collections import deque
from graphlib import CycleError, TopologicalSorter

from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple


class WorkbookFormulaGraph:
    def __init__(self, manifest: dict):
        self.manifest = manifest
        self.sheets = [s for s in manifest.get("sheets", []) if isinstance(s, dict)]
        self.sheet_names = {str(s.get("name", "")).casefold(): str(s.get("name", "")) for s in self.sheets}
        self.cells = {}
        self.by_sheet = {}
        self.row_labels = {}
        self.column_labels = {}
        self.names = {}
        self.tables = {}
        for sheet in self.sheets:
            name = str(sheet.get("name", ""))
            self.by_sheet[name] = []
            for cell in sheet.get("cells", []):
                address = str(cell.get("coordinate", "")).replace("$", "").upper()
                if not re.fullmatch(r"[A-Z]{1,3}[1-9]\d*", address):
                    continue
                self.cells[(name, address)] = cell
                self.by_sheet[name].append((address, cell))
                label = cell.get("cached_value") if self.formula(cell) else cell.get("value")
                if isinstance(label, str) and not label.startswith("="):
                    row, column = coordinate_to_tuple(address)
                    self.row_labels.setdefault((name, row), []).append((column, address, label))
                    self.column_labels.setdefault((name, column), []).append((row, address, label))
            for table in sheet.get("tables", []):
                self.tables[str(table.get("name", "")).casefold()] = (name, table)
            self._load_names(sheet.get("defined_names", []), name)
        self._load_names((manifest.get("workbook") or {}).get("defined_names", []), None)

    def _load_names(self, names, scope):
        for value in names:
            name_scope = scope
            if isinstance(value, dict):
                name = value.get("name")
                expression = value.get("attr_text") or value.get("value") or ""
                local_id = value.get("localSheetId")
                if local_id is not None and 0 <= int(local_id) < len(self.sheets):
                    name_scope = str(self.sheets[int(local_id)].get("name", ""))
            else:
                text = str(value)
                # Accept old serialized openpyxl names and simple name=range manifests.
                match = re.search(r"name='([^']+)'.*attr_text=(?:'([^']*)'|\"([^\"]*)\")", text, re.S)
                if match:
                    name, expression = match.group(1), match.group(2) or match.group(3)
                elif "=" in text and not text.startswith("<"):
                    name, expression = text.split("=", 1)
                else:
                    continue
            if name:
                self.names[(name_scope, str(name).casefold())] = str(expression)

    @staticmethod
    def formula(cell: dict) -> str:
        value = cell.get("value")
        if isinstance(value, str) and value.startswith("="):
            return value
        formula = cell.get("formula")
        return str(formula) if formula else ""

    def _range(self, sheet: str, address: str) -> tuple[list[tuple[str, str]], list[str]]:
        sheet = self.sheet_names.get(sheet.casefold(), sheet)
        if sheet not in self.by_sheet:
            return [], [f"Missing sheet: {sheet}"]
        address = address.replace("$", "").upper()
        try:
            col1, row1, col2, row2 = range_boundaries(address)
        except ValueError:
            return [], [f"Unsupported reference: {sheet}!{address}"]
        if col1 is None and row1 is None:
            return [], [f"Unsupported reference: {sheet}!{address}"]
        if ":" not in address and col1 and row1:
            key = (sheet, address)
            return ([key], []) if key in self.cells else ([], [f"Blank or absent cell in extraction: {sheet}!{address}"])
        found = []
        for coordinate, _ in self.by_sheet[sheet]:
            row, col = coordinate_to_tuple(coordinate)
            if ((col1 is None or col1 <= col <= col2) and (row1 is None or row1 <= row <= row2)):
                found.append((sheet, coordinate))
        return found, ([] if found else [f"No extracted cells in range: {sheet}!{address}"])

    def _operand(self, value: str, current_sheet: str, seen_names=frozenset()):
        if "!" in value:
            raw_sheet, address = value.rsplit("!", 1)
            sheet = raw_sheet.strip("'").replace("''", "'")
            if "[" in sheet or "]" in sheet:
                return [], [f"External workbook reference requires its source workbook: {value}"]
            if ":" in sheet:
                start, end = sheet.split(":", 1)
                names = list(self.by_sheet)
                start = self.sheet_names.get(start.casefold(), start)
                end = self.sheet_names.get(end.casefold(), end)
                if start not in names or end not in names:
                    return [], [f"Unresolved 3D reference: {value}"]
                first, last = names.index(start), names.index(end)
                if first > last:
                    return [], [f"Unresolved 3D reference: {value}"]
                refs, gaps = [], []
                for name in names[first:last + 1]:
                    cells, warnings = self._range(name, address)
                    refs.extend(cells)
                    gaps.extend(warnings)
                return refs, gaps
            current_sheet = self.sheet_names.get(sheet.casefold(), sheet)
            value = address
        name_key = (current_sheet, value.casefold())
        if name_key not in self.names:
            name_key = (None, value.casefold())
        if name_key in self.names:
            if name_key in seen_names:
                return [], [f"Circular defined name: {value}"]
            return self._expression(self.names[name_key], current_sheet, seen_names | {name_key})
        if "[" in value:
            return self._structured_reference(value, current_sheet)
        return self._range(current_sheet, value)

    def _structured_reference(self, value, current_sheet):
        name = value.split("[", 1)[0]
        entry = self.tables.get(name.casefold())
        if not entry:
            return [], [f"Unresolved structured table reference: {value}"]
        sheet, table = entry
        try:
            col1, row1, col2, row2 = range_boundaries(table.get("range") or table.get("ref"))
        except (TypeError, ValueError):
            return [], [f"Missing table range for {value}"]
        columns = table.get("columns") or []
        headers = {str(c.get("name", "")).casefold(): col1 + index for index, c in enumerate(columns)}
        if not headers:
            for address, cell in self.by_sheet[sheet]:
                row, col = coordinate_to_tuple(address)
                if row == row1 and col1 <= col <= col2:
                    headers[str(cell.get("value", "")).casefold()] = col
        names = [text for text in re.findall(r"\[([^\[\]]+)\]", value) if not text.startswith("#")]
        if any(text.startswith("@") for text in names) or "#This Row" in value:
            return [], [f"Row-relative table reference needs Excel row evaluation: {value}"]
        requested = [headers[text.casefold()] for text in names if text.casefold() in headers]
        if names and len(requested) != len(names):
            return [], [f"Missing table column metadata for {value}"]
        if "#All" not in value:
            if "#Headers" in value:
                row2 = row1
            elif "#Totals" in value:
                row1 = row2
            else:
                row1 += int(table.get("headerRowCount", 1))
                row2 -= int(table.get("totalsRowCount") or 0)
        found = []
        for address, _ in self.by_sheet[sheet]:
            row, col = coordinate_to_tuple(address)
            selected_column = min(requested) <= col <= max(requested) if requested else col1 <= col <= col2
            if row1 <= row <= row2 and selected_column:
                found.append((sheet, address))
        return found, []

    def _expression(self, formula: str, sheet: str, seen_names=frozenset()):
        formula = formula if formula.startswith("=") else "=" + formula
        try:
            tokens = Tokenizer(formula).items
        except Exception:
            return [], [f"Formula could not be parsed: {formula}"]
        refs, warnings = [], []
        for token in tokens:
            if token.type == "OPERAND" and token.subtype == "RANGE":
                found, gaps = self._operand(token.value, sheet, seen_names)
                refs.extend(found)
                warnings.extend(gaps)
            elif token.type == "OPERAND" and token.subtype == "ERROR":
                warnings.append(f"Excel error in formula: {token.value}")
            elif token.type == "FUNC" and token.subtype == "OPEN" and token.value.upper().removeprefix("_XLFN.") in {"INDIRECT(", "OFFSET("}:
                warnings.append(f"Dynamic reference requires Excel evaluation: {token.value[:-1]}")
        return list(dict.fromkeys(refs)), list(dict.fromkeys(warnings))

    def roots_in_range(self, metadata: dict) -> list[tuple[str, str]]:
        sheet = self.sheet_names.get(str(metadata.get("sheet_name", "")).casefold(), "")
        if not sheet:
            return []
        rows = (int(metadata.get("row_start") or 1), int(metadata.get("row_end") or 1048576))
        columns = range_boundaries(f"{metadata.get('column_start') or 'A'}1:{metadata.get('column_end') or 'XFD'}1")
        result = []
        for coordinate, cell in self.by_sheet[sheet]:
            row, col = coordinate_to_tuple(coordinate)
            if rows[0] <= row <= rows[1] and columns[0] <= col <= columns[2] and self.formula(cell):
                result.append((sheet, coordinate))
        return result

    def precedents(self, roots, *, max_cells=256, max_depth=6):
        """Bound a complete dependency walk explicitly; never silently imply completeness."""
        queue = deque((key, 0, ()) for key in dict.fromkeys(roots))
        root_set = set(roots)
        visited, dependencies, warnings, edges = set(), [], [], []
        while queue:
            key, depth, path = queue.popleft()
            if key in path:
                warnings.append(f"Circular reference: {key[0]}!{key[1]}")
                continue
            if key in visited:
                continue
            visited.add(key)
            cell = self.cells.get(key, {})
            if key not in root_set:
                if len(dependencies) >= max_cells:
                    warnings.append(f"Dependency context capped at {max_cells} cells; additional precedents omitted")
                    break
                dependencies.append(key)
            formula = self.formula(cell)
            if not formula:
                continue
            if cell.get("cached_value") is None:
                warnings.append(f"No saved Excel result for {key[0]}!{key[1]}; not recalculated")
            refs, gaps = self._expression(formula, key[0])
            warnings.extend(f"{key[0]}!{key[1]}: {gap}" for gap in gaps)
            edges.extend((key, ref) for ref in refs)
            if refs and depth >= max_depth:
                warnings.append(f"Dependency depth capped at {max_depth} for {key[0]}!{key[1]}")
                continue
            queue.extend((ref, depth + 1, (*path, key)) for ref in refs)
        adjacency = {}
        for source, target in edges:
            adjacency.setdefault(source, set()).add(target)
        try:
            tuple(TopologicalSorter(adjacency).static_order())
        except CycleError as exc:
            cycle = " -> ".join(f"{s}!{a}" for s, a in exc.args[1])
            warnings.append(f"Circular reference: {cycle}")
        return {"cells": dependencies, "warnings": list(dict.fromkeys(warnings)), "edges": edges}

    def render_cell(self, key) -> str:
        cell = self.cells[key]
        formula = self.formula(cell)
        value = cell.get("value")
        if formula:
            result = f"{key[1]}={formula} [saved Excel result: {cell.get('cached_value')!r}; not recalculated]"
        else:
            result = f"{key[1]}={value!r}"
        if cell.get("number_format"):
            result += f" [format: {cell['number_format']}]"
        # Include a row label and nearby column heading to make the input interpretable.
        row, column = coordinate_to_tuple(key[1])
        labels = []
        labels.extend((a, v) for c, a, v in self.row_labels.get((key[0], row), []) if c < column)
        column_labels = self.column_labels.get((key[0], column), [])
        labels.extend((a, v) for r, a, v in column_labels if 0 < row - r <= 3)
        # Financial periods usually sit at the top of a schedule, well beyond
        # three rows from EBITDA, D&A, tax and cash-flow values.
        periods = [(r, a, v) for r, a, v in column_labels if r < row and
                   re.fullmatch(r"(?:FY|CY)\s*\d{2,4}[AEF]?|20\d{2}(?:[-/]\d{2,4})?[AEF]?", v.strip(), re.I)]
        if periods:
            _, address, period = max(periods)
            labels.append((address, period))
        if labels:
            result += " [labels: " + "; ".join(f"{a}={v}" for a, v in labels[-4:]) + "]"
        return result

    def named_range_text(self) -> str:
        return "\n".join(
            f"{name}={value}" + (f" [scope: {scope}]" if scope else "")
            for (scope, name), value in self.names.items()
        )
