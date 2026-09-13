import io
from unittest.mock import Mock, patch

import fitz
from django.test import SimpleTestCase, override_settings
from openpyxl import Workbook

from ai_orchestrator.services.document_processor import DocumentProcessorService


class ChatExtractionTests(SimpleTestCase):
    def setUp(self):
        self.service = DocumentProcessorService()
        self.service.provider = Mock()

    @patch.object(DocumentProcessorService, "_remote_extract", side_effect=AssertionError("docproc called"))
    def test_text_upload_needs_neither_docproc_nor_inference(self, remote):
        result = self.service.get_chat_extraction_result(b"Revenue: 0", "note.txt")
        self.assertEqual(result["text"], "Revenue: 0")
        self.service.provider.execute_standard.assert_not_called()

    def test_native_pdf_is_read_without_vision(self):
        with fitz.open() as doc:
            doc.new_page().insert_text((72, 72), "Acme revenue 123")
            result = self.service.get_chat_extraction_result(doc.tobytes(), "report.pdf")
        self.assertIn("Acme revenue 123", result["text"])
        self.assertIn("PAGE 1", result["text"])
        self.service.provider.execute_standard.assert_not_called()

    @override_settings(ALLOW_SHARED_MODEL_DOCUMENT_VISION=True)
    @patch("ai_orchestrator.services.document_processor.AIRuntimeService.get_default_personality", return_value=None)
    @patch("ai_orchestrator.services.document_processor.AIRuntimeService.get_text_model", return_value="gemma-test")
    def test_mixed_pdf_uses_vision_only_for_scanned_page(self, model, personality):
        self.service.provider.execute_standard.return_value = {"response": "Scanned evidence"}
        with fitz.open() as doc:
            doc.new_page().insert_text((72, 72), "Native evidence")
            doc.new_page()
            result = self.service.get_chat_extraction_result(doc.tobytes(), "mixed.pdf")
        self.assertIn("Native evidence", result["text"])
        self.assertIn("Scanned evidence", result["text"])
        self.service.provider.execute_standard.assert_called_once()
        self.assertEqual(result["render_metadata"]["vision_pages"], 1)

    @patch("ai_orchestrator.services.document_processor.AIRuntimeService.get_default_personality", return_value=None)
    @patch("ai_orchestrator.services.document_processor.AIRuntimeService.get_text_model", return_value="gemma-test")
    @override_settings(ALLOW_SHARED_MODEL_DOCUMENT_VISION=True)
    def test_empty_vision_output_is_failure_not_page_marker(self, model, personality):
        self.service.provider.execute_standard.return_value = {"response": ""}
        result = self.service.get_chat_extraction_result(b"image", "scan.png")
        self.assertEqual(result["text"], "")
        self.assertEqual(result["transcription_status"], "failed")

    def test_scanned_page_never_reaches_shared_text_model_by_default(self):
        with fitz.open() as doc:
            doc.new_page().insert_text((72, 72), "Native evidence")
            doc.new_page()
            result = self.service.get_chat_extraction_result(doc.tobytes(), "mixed.pdf")

        self.assertIn("Native evidence", result["text"])
        self.assertEqual(result["transcription_status"], "partial")
        self.assertEqual(result["render_metadata"]["failed_pages"], [2])
        self.service.provider.execute_standard.assert_not_called()

    def test_spreadsheet_preserves_zero_and_false(self):
        workbook = Workbook()
        workbook.active.append(["Revenue", 0, False])
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        result = self.service.get_chat_extraction_result(output.getvalue(), "data.xlsx")
        self.assertIn("A1=Revenue\tB1=0\tC1=False", result["text"])
        self.assertEqual(result["transcription_status"], "complete")
        self.assertIn("backend_fallback_extraction", result["quality_flags"])
        self.assertNotIn("partial_extraction", result["quality_flags"])

    @override_settings(ALLOW_SHARED_MODEL_DOCUMENT_VISION=False)
    @patch("ai_orchestrator.services.document_processor.PipelineRegistryService.render_prompt_stage", return_value=(None, "OCR", None))
    def test_remote_outage_cannot_fall_back_to_shared_model(self, prompt):
        self.service.docproc_url = "http://docproc.invalid"
        with patch.object(self.service, "_remote_extract", return_value=None):
            result = self.service.get_extraction_result(b"image", "scan.png")
        self.assertEqual(result["transcription_status"], "failed")
        self.service.provider.execute_standard.assert_not_called()

    def test_partial_evidence_attempts_dedicated_ocr(self):
        self.service.docproc_url = "http://docproc.invalid"
        partial = {"text": "native page", "transcription_status": "partial"}
        complete = {"text": "native page and scanned page", "transcription_status": "complete"}
        with patch.object(self.service, "get_chat_extraction_result", return_value=partial), patch.object(self.service, "get_extraction_result", return_value=complete) as remote:
            self.assertEqual(self.service.get_evidence_extraction_result(b"pdf", "mixed.pdf"), complete)
        remote.assert_called_once_with(b"pdf", "mixed.pdf", page_limit=None, allow_local_fallback=False)

    def test_partial_evidence_survives_dedicated_ocr_failure(self):
        self.service.docproc_url = "http://docproc.invalid"
        partial = {"text": "native page", "transcription_status": "partial"}
        with patch.object(self.service, "get_chat_extraction_result", return_value=partial), patch.object(self.service, "get_extraction_result", return_value={"text": "", "transcription_status": "failed"}):
            self.assertEqual(self.service.get_evidence_extraction_result(b"pdf", "mixed.pdf"), partial)

    def test_evidence_respects_remote_opt_out(self):
        self.service.docproc_url = "http://docproc.invalid"
        partial = {"text": "native page", "transcription_status": "partial"}
        with patch.object(self.service, "get_chat_extraction_result", return_value=partial), patch.object(self.service, "get_extraction_result") as remote:
            self.assertEqual(self.service.get_evidence_extraction_result(b"pdf", "mixed.pdf", allow_remote_fallback=False), partial)
        remote.assert_not_called()

    def test_spreadsheet_includes_formulas_and_last_sheet(self):
        workbook = Workbook()
        workbook.active['A1'] = '=SUM(B1:B10)'
        workbook.create_sheet('Last sheet')['Z200'] = 'LAST_CELL'
        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        result = self.service.get_chat_extraction_result(output.getvalue(), 'formulas.xlsx')
        self.assertIn('A1==SUM(B1:B10) [cached value: unavailable]', result['text'])
        self.assertIn('[Sheet: Last sheet]', result['text'])
        self.assertIn('Z200=LAST_CELL', result['text'])
